import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_exact_excel():
    output_path = r"C:\Github_Projects\DS3D-HVAC\Unity_DS3D_Tables_Exact.xlsx"
    downloads_path = r"C:\Users\user\Downloads\Unity_DS3D_Tables_Exact.xlsx"
    downloads_final_path = r"C:\Users\user\Downloads\Unity_DS3D_Tables_Exact_Final.xlsx"
    
    wb = openpyxl.Workbook()

    # ------------------------------------------------------------------
    # STYLES DEFINITION
    # ------------------------------------------------------------------
    font_main_title = Font(name="Calibri", size=20, bold=True, color="000000")
    font_sub_title = Font(name="Calibri", size=16, bold=True, color="000000")
    font_col_header = Font(name="Calibri", size=11, bold=True, color="000000")
    font_data_regular = Font(name="Calibri", size=10, bold=False, color="000000")
    font_to_unity_title = Font(name="Calibri", size=12, bold=True, color="000000")

    # Fills
    fill_blue_banner = PatternFill(start_color="DDEBF7", end_color="DDEBF7", fill_type="solid") # Soft baby blue
    fill_green_banner = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid") # Soft sage green
    fill_gold_header = PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid") # Golden amber
    
    # Borders
    thin_border_side = Side(border_style="thin", color="000000")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    # Alignments
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # ------------------------------------------------------------------
    # TABLE DATA
    # ------------------------------------------------------------------
    left_rows = [
        (1, "Active_Fault_Status", "Active Fault Status", "Code", "Integer", 0, "If Fault is active: 1, Fault Inactive: 0"),
        (2, "Current_Step_Number", "Current Step Number", "Step", "Integer", 0, "1-based step index in procedure (0 = reset/unstarted)"),
        (3, "Total_Steps", "Total Steps in Procedure", "Steps", "Integer", 0, "Total steps for active fault"),
        (4, "Total_Wrong_Attempts", "Total Wrong Attempts Counter", "Count", "Integer", 0, "Total cumulative wrong attempts taken across entire fault"),
        (5, "Fault_101_Status", "Fault 101: No HVAC operations", "State", "Boolean", 0, "0: Inactive, 1: Active (14 steps)"),
        (6, "Fault_102_Status", "Fault 102: No ventilation", "State", "Boolean", 0, "0: Inactive, 1: Active (22 steps)"),
        (7, "Fault_103_Status", "Fault 103: Reduced ventilation", "State", "Boolean", 0, "0: Inactive, 1: Active (15 steps)"),
        (8, "Fault_104_Status", "Fault 104: Suction low and discharge low (C1)", "State", "Boolean", 0, "0: Inactive, 1: Active (15 steps)"),
        (9, "Fault_105_Status", "Fault 105: Suction low and discharge low (C2)", "State", "Boolean", 0, "0: Inactive, 1: Active (14 steps)"),
        (10, "Fault_106_Status", "Fault 106: High suction and high discharge", "State", "Boolean", 0, "0: Inactive, 1: Active (26 steps)"),
        (11, "Fault_107_Status", "Fault 107: Compressor not operating", "State", "Boolean", 0, "0: Inactive, 1: Active (40 steps)"),
        (12, "Fault_108_Status", "Fault 108: Reduced cooling", "State", "Boolean", 0, "0: Inactive, 1: Active (33 steps)"),
        (13, "Fault_109_Status", "Fault 109: No Heating", "State", "Boolean", 0, "0: Inactive, 1: Active (15 steps)"),
        (14, "Fault_110_Status", "Fault 110: Over Cooling", "State", "Boolean", 0, "0: Inactive, 1: Active (8 steps)"),
        (15, "Fault_111_Status", "Fault 111: Over Heating", "State", "Boolean", 0, "0: Inactive, 1: Active (7 steps)"),
        (16, "Fault_112_Status", "Fault 112: No emergency ventilation", "State", "Boolean", 0, "0: Inactive, 1: Active (14 steps)"),
        (17, "Fault_113_Status", "Fault 113: Smoke identified in saloon", "State", "Boolean", 0, "0: Inactive, 1: Active (13 steps)"),
        (18, "Fault_114_Status", "Fault 114: Compressor / Blower noisy", "State", "Boolean", 0, "0: Inactive, 1: Active (36 steps)"),
        (19, "Fault_115_Status", "Fault 115: Spare Fault Scenario", "State", "Boolean", 0, "0: Inactive, 1: Active (Reserved in Excel)"),
        (20, "Fault_116_Status", "Fault 116: Spare Fault Scenario", "State", "Boolean", 0, "0: Inactive, 1: Active (Reserved in Excel)"),
        (21, "", "Spare ", "", "Boolean", "", ""),
        (22, "", "Spare ", "", "Boolean", "", ""),
        (23, "", "Spare ", "", "Boolean", "", ""),
        (24, "", "Spare ", "", "Boolean", "", ""),
        (25, "", "Spare ", "", "Boolean", "", ""),
        (26, "", "Spare ", "", "Boolean", "", ""),
        (27, "", "Spare ", "", "Boolean", "", ""),
        (28, "", "Spare ", "", "Boolean", "", ""),
        (29, "", "Spare ", "", "Boolean", "", ""),
        (30, "", "Spare ", "", "Boolean", "", ""),
        (31, "", "Spare ", "", "Boolean", "", ""),
        (32, "", "Spare ", "", "Boolean", "", ""),
        (33, "", "Spare ", "", "Boolean", "", ""),
        (34, "Telemetry_Event", "Telemetry Event Name", "String", "String", "idle", "fault_started / step_completed / fault_completed")
    ]

    # Updated Right Table: Rows 23-26 now have the actual 3D physical sensor highlight variables!
    right_rows = [
        (1, "BlueArrowOut", "Fresh Air arrow", "Fresh Air blue arrow", "Integer", 5, "0-10 (Speed)"),
        (2, "Fan1", "Condensor Fan RedArrow Set 1", "Condensor Fan RedArrow Set 1", "Integer", 5, "0-10 (Speed)"),
        (3, "Fan2", "Evaporator Fan RedArrow Set 2", "Evaporator Fan RedArrow Set 2", "Integer", 5, "0-10 (Speed)"),
        (4, "GreenArrowIN1", "Fresh Air Damper Right", "Fresh Air Damper Right", "Integer", 5, "0-10 (Speed)"),
        (5, "GreenArrowIN2", "Fresh Air Damper left", "Fresh Air Damper left", "Integer", 5, "0-10 (Speed)"),
        (6, "YellowArrowIn", "Return Air Damper", "Return Air Damper", "Integer", 5, "0-10 (Speed)"),
        (7, "RedArrow_Inlet", "Fresh Air arrow Colour", "Intensity", "Float", 0, "0: Red, 1: Green, 2: Yellow, 3: Blue"),
        (8, "RedArrow_Outlet", "Condensor Fan RedArrow Set 1 Colour", "Intensity", "Float", 0, "0: Red, 1: Green, 2: Yellow, 3: Blue"),
        (9, "BlueArrow_Outlet", "Evaporator Fan RedArrow Set 2 Colour", "Intensity", "Float", 0, "0: Red, 1: Green, 2: Yellow, 3: Blue"),
        (10, "GreenArrow_Inlet", "Fresh Air Inlet Color Intensity", "Intensity", "Float", 0, "0-10 (Material color lerp)"),
        (11, "YellowArrow_Inlet", "Return Air Inlet Color Intensity", "Intensity", "Float", 0, "0-10 (Material color lerp)"),
        (12, "FADL", "Fresh Air Damper Left Position", "Position", "Float", 0, "0-10 (Mesh blendshape weight)"),
        (13, "FADR", "Fresh Air Damper Right Position", "Position", "Float", 0, "0-10 (Mesh blendshape weight)"),
        (14, "RAD", "Return Air Damper Position", "Position", "Float", 0, "0-10 (Mesh blendshape weight)"),
        (15, "DS3D_OFF", "DS3D Application Quit", "DS3D OFF", "Integer", 0, "0: Running, 1: Shutdown"),
        (16, "Fault_Code_Trigger", "Fault Trigger Scenario Code", "Code", "Integer", 104, "Active Fault to Unity (101-114 active)"),
        (17, "Fault_reset", "Fault Reset Trigger", "Command", "Integer", 0, "0: Idle, 1: Reset"),
        (18, "ShowHideObjects", "Toggle Component Visibility", "State", "Boolean", 0, "0: Hide, 1: Show"),
        (19, "Compresor 1_Highlights", "Toggle Compressor Highlights Visibility", "State", "Boolean", 1, "0: Hide Highlights, 1: Show Highlights"),
        (20, "Compresor 2_Highlights", "Toggle Compressor Highlights Visibility", "State", "Boolean", 1, "0: Hide Highlights, 1: Show Highlights"),
        (21, "Compresor 3_Highlights", "Toggle Compressor Highlights Visibility", "State", "Boolean", 1, "0: Hide Highlights, 1: Show Highlights"),
        (22, "Compresor 4_Highlights", "Toggle Compressor Highlights Visibility", "State", "Boolean", 1, "0: Hide Highlights, 1: Show Highlights"),
        (23, "FAT_Sensor_Highlights", "Toggle FAT Sensor Highlight Visibility", "State", "Boolean", 1, "0: Hide Highlights, 1: Show Highlights (Fresh Air Temp)"),
        (24, "RAT_Sensor_Highlights", "Toggle RAT Sensor Highlight Visibility", "State", "Boolean", 1, "0: Hide Highlights, 1: Show Highlights (Return Air Temp)"),
        (25, "SAT_Sensor_Highlights", "Toggle SAT Sensor Highlight Visibility", "State", "Boolean", 1, "0: Hide Highlights, 1: Show Highlights (Supply Air Temp)"),
        (26, "Smoke_Detector_Highlights", "Toggle Smoke Detector Highlight Visibility", "State", "Boolean", 1, "0: Hide Highlights, 1: Show Highlights (Smoke Detector)"),
        (27, "", "Spare", "", "Integer", 0, ""),
        (28, "", "Spare", "", "Integer", 0, ""),
        (29, "", "Spare", "", "Integer", 0, ""),
        (30, "", "Spare", "", "Integer", 0, ""),
        (31, "", "Spare", "", "Integer", 0, ""),
        (32, "", "Spare", "", "Integer", 0, ""),
        (33, "", "Spare", "", "Integer", 0, ""),
        (34, "", "Spare", "", "Integer", 0, ""),
        (35, "", "Spare", "", "Integer", 0, ""),
        (36, "", "Spare", "", "Integer", 0, ""),
        (37, "", "Spare", "", "Integer", 0, ""),
        (38, "", "Spare", "", "Integer", 0, ""),
        (39, "", "Spare", "", "Integer", 0, ""),
        (40, "", "Spare", "", "Integer", 0, ""),
        (41, "", "Spare", "", "Integer", 0, ""),
        (42, "", "Spare", "", "Integer", 0, ""),
        (43, "", "Spare", "", "Integer", 0, ""),
        (44, "", "Spare", "", "Integer", 0, ""),
        (45, "", "Spare", "", "Integer", 0, ""),
        (46, "", "Spare", "", "Integer", 0, ""),
        (47, "", "Spare", "", "Integer", 0, ""),
        (48, "", "Spare", "", "Integer", 0, ""),
        (49, "", "Spare", "", "Integer", 0, ""),
        (50, "", "Spare", "", "Integer", 0, "")
    ]

    to_unity_requirements = [
        ("Every Arrow with Speed", ""),
        ("Every Arrow Colour Control", ""),
        ("Application Quit", ""),
        ("Fresh Air Damper Left Position", ""),
        ("Fresh Air Damper Right Position", ""),
        ("Return Air Damper Position", ""),
        ("Fault Code", "Code to trigger the repective fault on Unity"),
        ("Fault Reset Trigger", "Reset Entire Screen to Normal Mode"),
        ("Show Hide Objects", ""),
        ("All Component Highlight", "As per discussion")
    ]

    headers_left = ["SR.No", "Variable", "Labels", "Units ", "Data Type", "Value", "Remark"]
    headers_right = ["SR.No", "Variable", "Labels", "Units ", "Data Type", "Value", "Remark"]

    def populate_sheet(ws, start_row):
        ws.views.sheetView[0].showGridLines = True

        # Main Title Banner: Unity_DS3D (CPU_WS) across Cols A to O
        title_row = start_row
        ws.merge_cells(start_row=title_row, start_column=1, end_row=title_row, end_column=15)
        cell_main = ws.cell(row=title_row, column=1, value="Unity_DS3D (CPU_WS)")
        cell_main.font = font_main_title
        cell_main.fill = fill_blue_banner
        cell_main.alignment = align_center
        ws.row_dimensions[title_row].height = 42

        for c in range(1, 16):
            ws.cell(row=title_row, column=c).border = cell_border
            ws.cell(row=title_row, column=c).fill = fill_blue_banner

        # Blank spacer rows
        ws.row_dimensions[title_row + 1].height = 14
        ws.row_dimensions[title_row + 2].height = 14

        # Sub-headers row: Unity INPUTS from DS3D (A:G), OUTPUT TO DS3D (I:O), TO UNITY (R:S)
        sub_row = title_row + 3
        ws.row_dimensions[sub_row].height = 36

        # Left Header
        ws.merge_cells(start_row=sub_row, start_column=1, end_row=sub_row, end_column=7)
        c_left_hdr = ws.cell(row=sub_row, column=1, value="Unity INPUTS from DS3D")
        c_left_hdr.font = font_sub_title
        c_left_hdr.fill = fill_green_banner
        c_left_hdr.alignment = align_center
        for c in range(1, 8):
            ws.cell(row=sub_row, column=c).border = cell_border
            ws.cell(row=sub_row, column=c).fill = fill_green_banner

        # Spacer Col H
        ws.cell(row=sub_row, column=8).value = ""

        # Right Header
        ws.merge_cells(start_row=sub_row, start_column=9, end_row=sub_row, end_column=15)
        c_rgt_hdr = ws.cell(row=sub_row, column=9, value="OUTPUT TO DS3D")
        c_rgt_hdr.font = font_sub_title
        c_rgt_hdr.fill = fill_green_banner
        c_rgt_hdr.alignment = align_center
        for c in range(9, 16):
            ws.cell(row=sub_row, column=c).border = cell_border
            ws.cell(row=sub_row, column=c).fill = fill_green_banner

        # TO UNITY Header (Columns R & S)
        ws.merge_cells(start_row=sub_row, start_column=18, end_row=sub_row, end_column=19)
        c_tu_hdr = ws.cell(row=sub_row, column=18, value="TO UNITY")
        c_tu_hdr.font = font_to_unity_title
        c_tu_hdr.alignment = align_center
        ws.cell(row=sub_row, column=18).border = cell_border
        ws.cell(row=sub_row, column=19).border = cell_border

        # Column Header Row
        col_hdr_row = sub_row + 1
        ws.row_dimensions[col_hdr_row].height = 26

        for idx, h_text in enumerate(headers_left, start=1):
            cell = ws.cell(row=col_hdr_row, column=idx, value=h_text)
            cell.font = font_col_header
            cell.fill = fill_gold_header
            cell.alignment = align_center
            cell.border = cell_border

        ws.cell(row=col_hdr_row, column=8).value = ""

        for idx, h_text in enumerate(headers_right, start=9):
            cell = ws.cell(row=col_hdr_row, column=idx, value=h_text)
            cell.font = font_col_header
            cell.fill = fill_gold_header
            cell.alignment = align_center
            cell.border = cell_border

        # TO UNITY Column Headers (Cols R & S)
        c_req = ws.cell(row=col_hdr_row, column=18, value="Requirement")
        c_req.font = font_col_header
        c_req.alignment = align_center
        c_req.border = cell_border

        c_rem = ws.cell(row=col_hdr_row, column=19, value="Remarks")
        c_rem.font = font_col_header
        c_rem.alignment = align_center
        c_rem.border = cell_border

        # Populate Data Rows
        data_start_row = col_hdr_row + 1
        max_rows = max(len(left_rows), len(right_rows))

        for offset in range(max_rows):
            r = data_start_row + offset
            ws.row_dimensions[r].height = 24

            # Left Table
            if offset < len(left_rows):
                item = left_rows[offset]
                for c_idx, val in enumerate(item, start=1):
                    cell = ws.cell(row=r, column=c_idx, value=val)
                    cell.font = font_data_regular
                    cell.border = cell_border
                    if c_idx in [1, 4, 5, 6]:
                        cell.alignment = align_center
                    else:
                        cell.alignment = align_left
            else:
                for c_idx in range(1, 8):
                    cell = ws.cell(row=r, column=c_idx, value="")
                    cell.border = cell_border

            # Spacer Col H
            ws.cell(row=r, column=8).value = ""

            # Right Table
            if offset < len(right_rows):
                item = right_rows[offset]
                for c_idx, val in enumerate(item, start=9):
                    cell = ws.cell(row=r, column=c_idx, value=val)
                    cell.font = font_data_regular
                    cell.border = cell_border
                    if c_idx in [9, 12, 13, 14]:
                        cell.alignment = align_center
                    else:
                        cell.alignment = align_left
            else:
                for c_idx in range(9, 16):
                    cell = ws.cell(row=r, column=c_idx, value="")
                    cell.border = cell_border

            # Populate TO UNITY items in Cols R & S
            if offset < len(to_unity_requirements):
                tu_req, tu_rem = to_unity_requirements[offset]
                c_tu_r = ws.cell(row=r, column=18, value=tu_req)
                c_tu_r.font = font_data_regular
                c_tu_r.border = cell_border
                c_tu_r.alignment = align_center

                c_tu_s = ws.cell(row=r, column=19, value=tu_rem)
                c_tu_s.font = font_data_regular
                c_tu_s.border = cell_border
                c_tu_s.alignment = align_center

        # Column Widths
        col_widths = {
            "A": 10.0,  # SR.No
            "B": 26.0,  # Variable
            "C": 44.0,  # Labels
            "D": 11.0,  # Units
            "E": 13.0,  # Data Type
            "F": 11.0,  # Value
            "G": 54.0,  # Remark
            "H": 6.0,   # Spacer
            "I": 10.0,  # SR.No
            "J": 32.0,  # Variable
            "K": 42.0,  # Labels
            "L": 32.0,  # Units
            "M": 13.0,  # Data Type
            "N": 11.0,  # Value
            "O": 48.0,  # Remark
            "P": 5.0,   # Spacer
            "Q": 5.0,   # Spacer
            "R": 34.0,  # TO UNITY Requirement
            "S": 48.0   # TO UNITY Remarks
        }
        for col_l, w in col_widths.items():
            ws.column_dimensions[col_l].width = w

    # Single primary sheet: DATA TRANSFER
    ws1 = wb.active
    ws1.title = "DATA TRANSFER"
    populate_sheet(ws1, start_row=2)

    # Secondary sheet: DATA TRANSFER (Row 816)
    ws2 = wb.create_sheet(title="DATA TRANSFER (Row 816)")
    populate_sheet(ws2, start_row=816)

    # Save to project workspace
    wb.save(output_path)
    print(f"Saved: {output_path}")

    # Try saving to Downloads
    try:
        wb.save(downloads_path)
        print(f"Saved: {downloads_path}")
    except PermissionError:
        print(f"Note: {downloads_path} was open in Excel. Saving to {downloads_final_path}...")
        wb.save(downloads_final_path)
        print(f"Saved: {downloads_final_path}")

if __name__ == "__main__":
    create_exact_excel()
