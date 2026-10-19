import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_excel():
    output_path = r"C:\Github_Projects\DS3D-HVAC\DS3D_HVAC_Unity_Interface_Specification.xlsx"
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # STYLES DEFINITION
    # -------------------------------------------------------------
    font_title = Font(name="Segoe UI", size=15, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="E0E6F0")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_cat = Font(name="Segoe UI", size=10, bold=True, color="0D233A")
    font_regular = Font(name="Segoe UI", size=9, color="000000")
    font_code = Font(name="Consolas", size=9, bold=True, color="1F4E79")
    font_bold = Font(name="Segoe UI", size=9, bold=True, color="000000")
    font_req = Font(name="Segoe UI", size=9, bold=True, color="274E13")

    fill_navy = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    fill_blue_hdr = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    fill_green_hdr = PatternFill(start_color="385723", end_color="385723", fill_type="solid")
    fill_purple_hdr = PatternFill(start_color="5B2C6F", end_color="5B2C6F", fill_type="solid")
    fill_cat_blue = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    fill_cat_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    fill_cat_purple = PatternFill(start_color="F2E6F7", end_color="F2E6F7", fill_type="solid")
    fill_zebra = PatternFill(start_color="F9FAFC", end_color="F9FAFC", fill_type="solid")
    fill_highlight = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="D9D9D9")
    thin_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    header_border = Border(
        left=Side(border_style="thin", color="FFFFFF"),
        right=Side(border_style="thin", color="FFFFFF"),
        top=Side(border_style="medium", color="1F4E79"),
        bottom=Side(border_style="medium", color="1F4E79")
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")

    # =============================================================
    # SHEET 1: Overview & Core Requirements
    # =============================================================
    ws1 = wb.active
    ws1.title = "Overview & Requirements"
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner
    ws1.merge_cells("A1:H1")
    ws1["A1"] = "DS3D-HVAC Simulation — Unity 3D & LabVIEW Interface Specification"
    ws1["A1"].font = font_title
    ws1["A1"].fill = fill_navy
    ws1["A1"].alignment = align_center
    ws1.row_dimensions[1].height = 36

    ws1.merge_cells("A2:H2")
    ws1["A2"] = "Clean, Minimalist Interface Specification between LabVIEW/Excel and Unity DS3D | Prepared for Polymator & Development Teams"
    ws1["A2"].font = font_subtitle
    ws1["A2"].fill = fill_blue_hdr
    ws1["A2"].alignment = align_center
    ws1.row_dimensions[2].height = 22

    # Executive Summary
    ws1["A4"] = "1. PURPOSE & ARCHITECTURAL SEPARATION"
    ws1["A4"].font = Font(name="Segoe UI", size=12, bold=True, color="1F4E79")
    
    summary_text = [
        "This workbook replaces the confusing 1,500-row master Excel by isolating ONLY the variables that Unity 3D and LabVIEW exchange.",
        "• LabVIEW Internal Domain: Mathematical calculations (passenger heat load in Watts, solar gain, sensible heat delta T, 27 circuit relays, TCMS cab switches) stay inside LabVIEW.",
        "• Unity 3D Domain: 3D interactive procedural training, tool interactions, mesh blendshapes (damper positions), animated airflow arrows, and component highlights.",
        "• Network Protocol: Real-time UDP socket over localhost (or LAN IP) — LabVIEW sends controls to Port 9000; Unity sends telemetry back to Port 5005."
    ]
    for idx, text in enumerate(summary_text, start=5):
        ws1.merge_cells(f"A{idx}:H{idx}")
        ws1[f"A{idx}"] = text
        ws1[f"A{idx}"].font = font_regular
        ws1.row_dimensions[idx].height = 18

    # The 10 Core Requirements Table
    ws1["A10"] = "2. THE 10 CORE UNITY REQUIREMENTS (As requested by Excel Sir)"
    ws1["A10"].font = Font(name="Segoe UI", size=12, bold=True, color="1F4E79")

    req_headers = ["Req #", "Requirement Name", "Direction", "Excel/LabVIEW Variables Mapped", "Data Type", "Function in Unity 3D", "Remarks & Clarifications"]
    for col_idx, h in enumerate(req_headers, start=1):
        cell = ws1.cell(11, col_idx, h)
        cell.font = font_header
        cell.fill = fill_blue_hdr
        cell.alignment = align_center
        cell.border = header_border
    ws1.row_dimensions[11].height = 26

    req_data = [
        (1, "Every Arrow with Speed", "LabVIEW → Unity", "BlueArrowOut, Fan1, Fan2, GreenArrowIN1, GreenArrowIN2, YellowArrowIn", "Integer / Float (0-10)", "Controls playback speed of airflow arrows in 3D scene", "Default = 5. Matches 6 distinct airflow arrow animation sets."),
        (2, "Every Arrow Colour Control", "LabVIEW → Unity", "RedArrow_Inlet, RedArrow_Outlet, BlueArrow_Outlet, GreenArrow_Inlet, YellowArrow_Inlet", "Float (0-10)", "Lerps material color intensity between cold/ambient/hot states", "Default = 0. Controls 5 arrow color channels via material shader."),
        (3, "Application Quit", "LabVIEW → Unity", "DS3D_OFF", "Integer (0/1)", "Terminates the Unity standalone application cleanly", "0: Running, 1: Shutdown. Sent when session or exam closes."),
        (4, "Fresh Air Damper Left Position", "LabVIEW → Unity", "FADL", "Float (0-10 / 0-100%)", "Controls Left Damper louver mechanical opening via Mesh BlendShape", "Physically rotates damper blades from 0% (closed) to 100% (open)."),
        (5, "Fresh Air Damper Right Position", "LabVIEW → Unity", "FADR", "Float (0-10 / 0-100%)", "Controls Right Damper louver mechanical opening via Mesh BlendShape", "Physically rotates damper blades from 0% (closed) to 100% (open)."),
        (6, "Return Air Damper Position", "LabVIEW → Unity", "RAD", "Float (0-10 / 0-100%)", "Controls Return Air Damper mechanical opening via Mesh BlendShape", "Normal = Open (100%). In Emergency / Smoke mode = Closed (0%)."),
        (7, "Fault Code", "LabVIEW → Unity", "Fault_Code_Trigger, Fault_Cause_Trigger", "Integer (101-114)", "Triggers the active procedural fault training scenario in Unity", "Code 101 to 114 triggers the respective procedural overhaul sequence."),
        (8, "Fault Reset Trigger", "LabVIEW → Unity", "Fault_Command_Action", "Integer (0, 1, 2)", "Resets entire 3D simulation to normal idle operating mode", "0: Idle, 1: Trigger active fault, 2: Reset entire screen/scene."),
        (9, "Show Hide Objects", "LabVIEW → Unity", "ShowHideObjects", "Boolean (0/1)", "Toggles outer sheet-metal casing / protective covers on/off", "0: Hide casing (see inside internal components), 1: Show casing."),
        (10, "All Component Highlight", "LabVIEW → Unity", "Compressor 1-4, FAD L/R, RAD, All Sensors, Fans, Valves", "Boolean (0/1 per part)", "Activates visual outline / highlight glow on target components", "Enables technician/student to visually identify target parts for inspection.")
    ]

    for row_idx, r in enumerate(req_data, start=12):
        ws1.row_dimensions[row_idx].height = 24
        fill = fill_zebra if row_idx % 2 == 0 else PatternFill(fill_type=None)
        for col_idx, val in enumerate(r, start=1):
            cell = ws1.cell(row_idx, col_idx, val)
            cell.font = font_code if col_idx in [1, 4] else font_regular
            cell.border = thin_border
            cell.fill = fill
            cell.alignment = align_center if col_idx in [1, 3, 5] else align_left

    # Explanation of doubts
    r_start = 24
    ws1[f"A{r_start}"] = "3. DIRECT ANSWERS TO YOUR 4 DOUBTS"
    ws1[f"A{r_start}"].font = Font(name="Segoe UI", size=12, bold=True, color="1F4E79")

    doubts_text = [
        ("Doubt 1: Fresh Air Damper needs only highlights, no animation right?",
         "NO — Fresh Air Dampers have BOTH physical vane opening animation AND highlights. The physical opening (0% to 100%) is driven by mesh BlendShapes (FADL and FADR). Highlights are outlines used during repair procedures."),
        ("Doubt 2: What is Return Air Damper Position?",
         "RAD Position is the opening/closing angle of the Return Air Damper blades (RAD). Normally open (100%) for recirculating saloon air; closed (0%) during Smoke or Fire modes so smoke is not blown into passengers."),
        ("Doubt 3: What is show hide objects?",
         "Show/Hide Objects (ShowHideObjects) toggles the HVAC outer sheet-metal casing/cover. When Hidden (0), the technician can see inside the unit to inspect compressors and piping. When Shown (1), the cover is restored."),
        ("Doubt 4: All component Highlights include: 2 Fresh Air Dampers, All sensors, all compressors right?",
         "YES, EXACTLY. It covers 4 Compressors, 2 Fresh Air Dampers (Left & Right), 1 Return Air Damper, and all 7 Sensors (RAT, SAT, FAT, Smoke Detector, HP Transducer, LP Transducer, Humidity Sensor), plus contactors & valves.")
    ]

    curr_r = r_start + 1
    for title, desc in doubts_text:
        ws1.merge_cells(f"A{curr_r}:H{curr_r}")
        ws1[f"A{curr_r}"] = title
        ws1[f"A{curr_r}"].font = font_bold
        ws1[f"A{curr_r}"].fill = fill_cat_blue
        ws1.row_dimensions[curr_r].height = 20
        curr_r += 1

        ws1.merge_cells(f"A{curr_r}:H{curr_r}")
        ws1[f"A{curr_r}"] = desc
        ws1[f"A{curr_r}"].font = font_regular
        ws1.row_dimensions[curr_r].height = 22
        curr_r += 1

    # =============================================================
    # SHEET 2: LabVIEW_to_Unity (Controls)
    # =============================================================
    ws2 = wb.create_sheet(title="LabVIEW_to_Unity (Controls)")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:H1")
    ws2["A1"] = "OUTPUT TO DS3D — Controls Sent from LabVIEW / Excel to Unity (UDP Port 9000)"
    ws2["A1"].font = font_title
    ws2["A1"].fill = fill_navy
    ws2["A1"].alignment = align_center
    ws2.row_dimensions[1].height = 36

    ws2.merge_cells("A2:H2")
    ws2["A2"] = "Streamlined from the previous master table into 6 logical categories matching Excel Sir's 10 Core Requirements"
    ws2["A2"].font = font_subtitle
    ws2["A2"].fill = fill_blue_hdr
    ws2["A2"].alignment = align_center
    ws2.row_dimensions[2].height = 20

    headers_l2u = ["SR.No", "Variable Name", "Display Label", "Data Type", "Value Range", "Default", "Req Ref", "Detailed Description & Unity Action"]
    for col_idx, h in enumerate(headers_l2u, start=1):
        cell = ws2.cell(4, col_idx, h)
        cell.font = font_header
        cell.fill = fill_blue_hdr
        cell.alignment = align_center
        cell.border = header_border
    ws2.row_dimensions[4].height = 26

    # Data items for LabVIEW to Unity
    controls_data = [
        # Category 1: Arrow Speeds (Req 1)
        ("CATEGORY 1: AIRFLOW ARROW ANIMATION SPEEDS", [
            (1, "BlueArrowOut", "Fresh Air Blue Arrow Speed", "Integer", "0 - 10", 5, "Req 1", "Controls velocity of blue fresh air intake arrows entering the unit"),
            (2, "Fan1", "Condenser Fan 1 Red Arrow Speed", "Integer", "0 - 10", 5, "Req 1", "Controls velocity of red heat discharge arrows on Condenser Fan 1"),
            (3, "Fan2", "Evaporator Fan 2 Red Arrow Speed", "Integer", "0 - 10", 5, "Req 1", "Controls velocity of airflow arrows exiting Evaporator Blower 2"),
            (4, "GreenArrowIN1", "Fresh Air Damper Right Arrow Speed", "Integer", "0 - 10", 5, "Req 1", "Controls velocity of green inlet airflow arrows through Right FAD"),
            (5, "GreenArrowIN2", "Fresh Air Damper Left Arrow Speed", "Integer", "0 - 10", 5, "Req 1", "Controls velocity of green inlet airflow arrows through Left FAD"),
            (6, "YellowArrowIn", "Return Air Damper Arrow Speed", "Integer", "0 - 10", 5, "Req 1", "Controls velocity of yellow recirculation arrows through RAD")
        ]),
        # Category 2: Arrow Color Controls (Req 2)
        ("CATEGORY 2: AIRFLOW ARROW COLOR TEMPERATURE CONTROLS", [
            (7, "RedArrow_Inlet", "Hot Air Inlet Arrow Color", "Float", "0 - 10", 0, "Req 2", "Material color lerp: adjusts intensity/hue of inlet warm air stream"),
            (8, "RedArrow_Outlet", "Hot Air Outlet Arrow Color", "Float", "0 - 10", 0, "Req 2", "Material color lerp: adjusts intensity/hue of condenser discharge air"),
            (9, "BlueArrow_Outlet", "Cool Air Outlet Arrow Color", "Float", "0 - 10", 0, "Req 2", "Material color lerp: adjusts intensity/hue of cooled supply air stream"),
            (10, "GreenArrow_Inlet", "Fresh Air Inlet Color Intensity", "Float", "0 - 10", 0, "Req 2", "Material color lerp: adjusts visual prominence of outside fresh air"),
            (11, "YellowArrow_Inlet", "Return Air Inlet Color Intensity", "Float", "0 - 10", 0, "Req 2", "Material color lerp: adjusts visual prominence of saloon return air")
        ]),
        # Category 3: Damper Physical Positions (Req 4, 5, 6)
        ("CATEGORY 3: DAMPER PHYSICAL POSITIONS (MESH BLENDSHAPES)", [
            (12, "FADL", "Fresh Air Damper Left Position", "Float", "0 - 10 (0-100%)", 0, "Req 4", "Controls physical louver rotation of Left Fresh Air Damper (Mesh BlendShape)"),
            (13, "FADR", "Fresh Air Damper Right Position", "Float", "0 - 10 (0-100%)", 0, "Req 5", "Controls physical louver rotation of Right Fresh Air Damper (Mesh BlendShape)"),
            (14, "RAD", "Return Air Damper Position", "Float", "0 - 10 (0-100%)", 10, "Req 6", "Controls physical louver rotation of Return Air Damper (Normal=10/100%, Smoke=0%)")
        ]),
        # Category 4: Application & Visual Controls (Req 3, 9)
        ("CATEGORY 4: APPLICATION & 3D CASING CONTROLS", [
            (15, "DS3D_OFF", "DS3D Application Quit", "Integer", "0 or 1", 0, "Req 3", "0 = Application active, 1 = Application Quit / Close simulation"),
            (16, "ShowHideObjects", "Toggle Outer Casing Visibility", "Boolean", "0 or 1", 0, "Req 9", "0 = Hide outer covers to see internal components; 1 = Show outer casing")
        ]),
        # Category 5: Fault Scenario Commands (Req 7, 8)
        ("CATEGORY 5: FAULT SCENARIO COMMANDS", [
            (17, "Fault_Code_Trigger", "Active Fault Scenario Code", "Integer", "0, 101 - 114", 0, "Req 7", "0 = Idle/Normal. Code 101-114 activates the respective overhaul procedure"),
            (18, "Fault_Cause_Trigger", "Fault Cause Number", "Integer", "1 - 5", 1, "Req 7", "Selects the root-cause variation within the active fault (Cause 1, Cause 2, etc.)"),
            (19, "Fault_Command_Action", "Fault Command Action", "Integer", "0, 1, 2", 0, "Req 8", "0 = Idle, 1 = Trigger Fault, 2 = Reset Entire Screen / Simulator to Normal")
        ]),
        # Category 6: Component Highlights (Req 10)
        ("CATEGORY 6: COMPONENT HIGHLIGHTS (OUTLINE / GLOW)", [
            (20, "Compressor_1_Highlight", "Compressor 1 Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles yellow/cyan outline glow on Compressor 1 assembly"),
            (21, "Compressor_2_Highlight", "Compressor 2 Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles yellow/cyan outline glow on Compressor 2 assembly"),
            (22, "Compressor_3_Highlight", "Compressor 3 Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles yellow/cyan outline glow on Compressor 3 assembly"),
            (23, "Compressor_4_Highlight", "Compressor 4 Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles yellow/cyan outline glow on Compressor 4 assembly"),
            (24, "FAD_Left_Highlight", "Fresh Air Damper Left Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Fresh Air Damper Left assembly & motor"),
            (25, "FAD_Right_Highlight", "Fresh Air Damper Right Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Fresh Air Damper Right assembly & motor"),
            (26, "RAD_Damper_Highlight", "Return Air Damper Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Return Air Damper assembly & actuator"),
            (27, "RAT_Sensor_Highlight", "RAT Sensor Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Return Air Temperature Sensor probe"),
            (28, "SAT_Sensor_Highlight", "SAT Sensor Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Supply Air Temperature Sensor probe"),
            (29, "FAT_Sensor_Highlight", "FAT Sensor Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Fresh Air Temperature Sensor probe"),
            (30, "Smoke_Detector_Highlight", "Smoke Detector Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Internal Smoke Detector unit"),
            (31, "HP_Sensor_Highlight", "High Pressure Sensor Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on High Pressure Transducers/Switches"),
            (32, "LP_Sensor_Highlight", "Low Pressure Sensor Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Low Pressure Transducers/Switches"),
            (33, "Humidity_Sensor_Highlight", "Humidity Sensor Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Saloon Relative Humidity Sensor"),
            (34, "Condenser_Fan_1_Highlight", "Condenser Fan 1 Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Condenser Fan 1 motor & impellers"),
            (35, "Condenser_Fan_2_Highlight", "Condenser Fan 2 Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Condenser Fan 2 motor & impellers"),
            (36, "Supply_Air_Blower_1_Highlight", "Supply Air Blower 1 Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Evaporator Supply Air Blower 1"),
            (37, "Supply_Air_Blower_2_Highlight", "Supply Air Blower 2 Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Evaporator Supply Air Blower 2"),
            (38, "Filter_Drier_Highlight", "Filter Drier Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Refrigerant Filter Drier cartridge"),
            (39, "TXV_Valve_Highlight", "TXV Expansion Valve Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Thermal Expansion Valve (TXV) body"),
            (40, "Solenoid_Valve_Highlight", "Solenoid Valve Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Liquid Line Solenoid Valve"),
            (41, "Heater_Bank_Highlight", "Heater Bank Assembly Highlight", "Boolean", "0 or 1", 0, "Req 10", "Toggles outline glow on Electric Overhead Heating Bank")
        ])
    ]

    curr_row = 5
    for cat_title, items in controls_data:
        # Category header row
        ws2.merge_cells(f"A{curr_row}:H{curr_row}")
        ws2[f"A{curr_row}"] = cat_title
        ws2[f"A{curr_row}"].font = font_cat
        ws2[f"A{curr_row}"].fill = fill_cat_blue
        ws2.row_dimensions[curr_row].height = 22
        curr_row += 1

        for item in items:
            ws2.row_dimensions[curr_row].height = 20
            fill = fill_zebra if curr_row % 2 == 0 else PatternFill(fill_type=None)
            for c_idx, val in enumerate(item, start=1):
                cell = ws2.cell(curr_row, c_idx, val)
                cell.font = font_code if c_idx == 2 else (font_bold if c_idx in [1, 7] else font_regular)
                cell.border = thin_border
                cell.fill = fill
                cell.alignment = align_center if c_idx in [1, 4, 5, 6, 7] else align_left
            curr_row += 1

    # =============================================================
    # SHEET 3: Unity_to_LabVIEW (Telemetry)
    # =============================================================
    ws3 = wb.create_sheet(title="Unity_to_LabVIEW (Telemetry)")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:H1")
    ws3["A1"] = "INPUTS FROM DS3D — Telemetry Sent from Unity to LabVIEW / Excel (UDP Port 5005)"
    ws3["A1"].font = font_title
    ws3["A1"].fill = fill_navy
    ws3["A1"].alignment = align_center
    ws3.row_dimensions[1].height = 36

    ws3.merge_cells("A2:H2")
    ws3["A2"] = "Procedure progress, step accomplishments, mistake penalties, and scoring metrics for student evaluation"
    ws3["A2"].font = font_subtitle
    ws3["A2"].fill = fill_green_hdr
    ws3["A2"].alignment = align_center
    ws3.row_dimensions[2].height = 20

    headers_u2l = ["SR.No", "Variable Name", "Display Label", "Data Type", "Value Range", "Default", "Purpose in LabVIEW", "Why This Variable is Essential (Explanation)"]
    for col_idx, h in enumerate(headers_u2l, start=1):
        cell = ws3.cell(4, col_idx, h)
        cell.font = font_header
        cell.fill = fill_green_hdr
        cell.alignment = align_center
        cell.border = header_border
    ws3.row_dimensions[4].height = 26

    telemetry_data = [
        # Category 1: Live Step Progress & Scoring (Essential for Exam/Training Evaluation)
        ("CATEGORY 1: LIVE PROCEDURAL TELEMETRY & STUDENT SCORING", [
            (1, "Active_Fault_Code", "Active Fault Scenario Code", "Integer", "0, 101 - 114", 0, "Identifies current running fault on LabVIEW UI", "Essential so LabVIEW displays correct fault title and step instructions."),
            (2, "Current_Step_Number", "Current Step Number", "Integer", "1 - Total_Steps", 0, "Advances the step progress bar in LabVIEW", "Essential so instructor and student see live which step (e.g. Step 3) is active."),
            (3, "Total_Steps", "Total Steps in Procedure", "Integer", "4 - 40", 0, "Denominator for percentage completion calculation", "Essential to calculate Progress % = (Current_Step / Total_Steps) * 100."),
            (4, "Wrong_Attempts", "Wrong Tool Attempts Counter", "Integer", "0, 1, 2...", 0, "Increments when student uses incorrect tool", "CRITICAL FOR SCORING: Deducts penalty points per mistake in exam mode."),
            (5, "Step_Accomplished", "Step Accomplished Pulse", "Boolean", "0 or 1", 0, "Fires pulse (1) when step action completes", "Tells LabVIEW that step animation finished; triggers next step countdown."),
            (6, "Fault_Completed", "Fault Procedure Finished Flag", "Boolean", "0 or 1", 0, "Latches to 1 when final step is completed", "Triggers final test pass/fail calculation and assessment report generation."),
            (7, "Fault_Submit_Status", "Repaired Fault Submit Status", "Boolean", "0 or 1", 0, "Latches to 1 when student clicks Submit UI button", "Confirms that student officially submitted repair work before timer elapsed."),
            (8, "Steps_Attempts_Summary", "Step Attempts Summary Payload", "String", "Array e.g. [0,2,1,0]", "0", "Detailed mistake history across every step", "Provides granular audit trail for instructor station post-exam analytics."),
            (9, "Telemetry_Event", "Telemetry Event Name", "String", "idle / step_completed / fault_completed", "idle", "Standard event tag for LabVIEW JSON parser", "Used by LabVIEW Unflatten From JSON primitive to route data packets.")
        ]),
        # Category 2: Discrete Fault Status Flags (For Direct Boolean Array Wiring in LabVIEW)
        ("CATEGORY 2: DISCRETE FAULT STATUS BITS (101 - 114)", [
            (10, "Fault_101_Status", "Fault 101: No HVAC Operations", "Boolean", "0 or 1", 0, "Fault 101 active state bit", "Direct boolean LED indicator on LabVIEW instructor panel."),
            (11, "Fault_102_Status", "Fault 102: No Ventilation", "Boolean", "0 or 1", 0, "Fault 102 active state bit", "Direct boolean LED indicator on LabVIEW instructor panel."),
            (12, "Fault_103_Status", "Fault 103: Reduced Ventilation", "Boolean", "0 or 1", 0, "Fault 103 active state bit", "Direct boolean LED indicator on LabVIEW instructor panel."),
            (13, "Fault_104_Status", "Fault 104: Low Suction & Discharge (C1)", "Boolean", "0 or 1", 0, "Fault 104 active state bit", "Direct boolean LED indicator on LabVIEW instructor panel."),
            (14, "Fault_105_Status", "Fault 105: Low Suction & Discharge (C2)", "Boolean", "0 or 1", 0, "Fault 105 active state bit", "Direct boolean LED indicator on LabVIEW instructor panel."),
            (15, "Fault_106_Status", "Fault 106: High Suction & Discharge", "Boolean", "0 or 1", 0, "Fault 106 active state bit", "Direct boolean LED indicator on LabVIEW instructor panel."),
            (16, "Fault_107_Status", "Fault 107: Compressor Not Operating", "Boolean", "0 or 1", 0, "Fault 107 active state bit", "Direct boolean LED indicator on LabVIEW instructor panel."),
            (17, "Fault_108_Status", "Fault 108: Reduced Cooling", "Boolean", "0 or 1", 0, "Fault 108 active state bit", "Direct boolean LED indicator on LabVIEW instructor panel."),
            (18, "Fault_109_Status", "Fault 109: No Heating", "Boolean", "0 or 1", 0, "Fault 109 active state bit", "Direct boolean LED indicator on LabVIEW instructor panel."),
            (19, "Fault_110_Status", "Fault 110: Over Cooling", "Boolean", "0 or 1", 0, "Fault 110 active state bit", "Direct boolean LED indicator on LabVIEW instructor panel."),
            (20, "Fault_111_Status", "Fault 111: Over Heating", "Boolean", "0 or 1", 0, "Fault 111 active state bit", "Direct boolean LED indicator on LabVIEW instructor panel."),
            (21, "Fault_112_Status", "Fault 112: No Emergency Ventilation", "Boolean", "0 or 1", 0, "Fault 112 active state bit", "Direct boolean LED indicator on LabVIEW instructor panel."),
            (22, "Fault_113_Status", "Fault 113: Smoke Identified in Saloon", "Boolean", "0 or 1", 0, "Fault 113 active state bit", "Direct boolean LED indicator on LabVIEW instructor panel."),
            (23, "Fault_114_Status", "Fault 114: Compressor / Blower Noisy", "Boolean", "0 or 1", 0, "Fault 114 active state bit", "Direct boolean LED indicator on LabVIEW instructor panel.")
        ])
    ]

    curr_row = 5
    for cat_title, items in telemetry_data:
        ws3.merge_cells(f"A{curr_row}:H{curr_row}")
        ws3[f"A{curr_row}"] = cat_title
        ws3[f"A{curr_row}"].font = font_cat
        ws3[f"A{curr_row}"].fill = fill_cat_green
        ws3.row_dimensions[curr_row].height = 22
        curr_row += 1

        for item in items:
            ws3.row_dimensions[curr_row].height = 20
            fill = fill_zebra if curr_row % 2 == 0 else PatternFill(fill_type=None)
            for c_idx, val in enumerate(item, start=1):
                cell = ws3.cell(curr_row, c_idx, val)
                cell.font = font_code if c_idx == 2 else (font_bold if c_idx in [1, 4] else font_regular)
                cell.border = thin_border
                cell.fill = fill
                cell.alignment = align_center if c_idx in [1, 4, 5, 6] else align_left
            curr_row += 1

    # =============================================================
    # SHEET 4: Fault_Catalog
    # =============================================================
    ws4 = wb.create_sheet(title="Fault_Catalog")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells("A1:G1")
    ws4["A1"] = "HVAC Procedural Fault Training Catalog (Fault Codes 101 - 114)"
    ws4["A1"].font = font_title
    ws4["A1"].fill = fill_navy
    ws4["A1"].alignment = align_center
    ws4.row_dimensions[1].height = 36

    ws4.merge_cells("A2:G2")
    ws4["A2"] = "Complete registry of train HVAC maintenance & overhaul scenarios implemented in Unity 3D"
    ws4["A2"].font = font_subtitle
    ws4["A2"].fill = fill_purple_hdr
    ws4["A2"].alignment = align_center
    ws4.row_dimensions[2].height = 20

    headers_fc = ["Fault Code", "Fault Name / Scenario Title", "Subsystem Scope", "Steps Count", "Typical Root Cause", "Key Components Involved", "Primary Maintenance Tools"]
    for col_idx, h in enumerate(headers_fc, start=1):
        cell = ws4.cell(4, col_idx, h)
        cell.font = font_header
        cell.fill = fill_purple_hdr
        cell.alignment = align_center
        cell.border = header_border
    ws4.row_dimensions[4].height = 26

    fault_catalog_data = [
        (101, "No HVAC Operation", "Electrical & Power Distribution", 14, "Tripped circuit breakers / blown fuse", "Electrical Control Panel, Circuit Breakers, Power Plugs", "Spanner, Multimeter, Screwdriver, Insulated Gloves"),
        (102, "No Ventilation", "Air Supply & Blower Assembly", 22, "Defective air blower motor / jammed impeller", "Supply Air Blower, Blower Motor, Hinge Pins", "Spanner, Pliers, Screwdriver, Compressed Air Gun"),
        (103, "Reduced Ventilation", "Fresh Air Damper & Air Filter", 15, "Clogged filter / jammed fresh air damper motor", "Fresh Air Damper, Damper Motor, Air Filters", "Spanner, Screwdriver, Bristle Brush, Visual Inspection"),
        (104, "Suction Low & Discharge Low (C1)", "Refrigeration Circuit 1", 15, "Refrigerant leak / defective TX valve / damaged coil", "Evaporator Coil, TX Valve, Refrigerant Lines", "Electronic Leak Detector, Manifold Gauge, Spanner"),
        (105, "Suction Low & Discharge Low (C2)", "Refrigeration Circuit 2", 14, "Moisture contamination / clogged filter drier", "Filter Drier, Solenoid Valve, Sight Glass", "Lifting Device, Insulated Gloves, Screwdriver, Spanner"),
        (106, "High Suction & High Discharge", "Condenser & Heat Dissipation", 26, "Condenser fan failure / bent condenser fins", "Condenser Fan Motor, Condenser Coil Fins", "Fin Comb, Bristle Brush, Compressed Air, Spanner"),
        (107, "Compressor Not Operating", "Refrigerant Compressor Units", 40, "Compressor motor contactor failure / locked rotor", "Compressor Units 1-4, Compressor Motor Contactors", "Multimeter, Spanner, Lifting Device, Insulated Gloves"),
        (108, "Reduced Cooling", "Cooling Modulation & Sensors", 33, "Defective RAT/SAT temperature sensor / low airflow", "Return Air Temp Sensor, Supply Air Temp Sensor", "Clamp Thermometer, Calibrated Sensor, Spanner"),
        (109, "No Heating", "Overhead Electric Heating", 15, "Tripped heater thermostat / burnt heating element", "Electric Heater Bank, Heater Thermostats (1 & 2)", "Multimeter, Screwdriver, Clamp Thermometer, Spanner"),
        (110, "Over Cooling", "Thermal Control & Regulation", 8, "Stuck thermostat / faulty temperature feedback", "Return Air Temperature Sensor, Controller PCB", "Clamp Thermometer, Visual Inspection, Spanner"),
        (111, "Over Heating", "Thermal Protection & Limiters", 7, "Faulty high-limit thermostat / blower failure", "Heater High-Limit Thermostat, Air Blower", "Clamp Thermometer, Multimeter, Screwdriver"),
        (112, "No Emergency Ventilation", "Emergency DC Inverter & Blower", 14, "Emergency inverter contactor trip / DC supply fault", "Emergency Inverter, DC Contactor, Battery Supply", "Multimeter, Screwdriver, Insulated Gloves"),
        (113, "Smoke Identified in Saloon", "Fire Detection & Damper Isolation", 13, "Defective smoke detector head / faulty FADR relay", "Internal Smoke Detector, FAD Dampers, RAD Damper", "Bristle Brush, Spanner, Vacuum Cleaner"),
        (114, "Compressor / Blower Excessively Noisy", "Mechanical Vibration & Bearings", 36, "Worn motor bearings / loose compressor mounting bolts", "Compressor Mountings, Blower Bearings, Damper Seals", "Spanner, Torque Wrench, Visual Inspection, Gauge Set")
    ]

    for row_idx, r in enumerate(fault_catalog_data, start=5):
        ws4.row_dimensions[row_idx].height = 22
        fill = fill_zebra if row_idx % 2 == 0 else PatternFill(fill_type=None)
        for col_idx, val in enumerate(r, start=1):
            cell = ws4.cell(row_idx, col_idx, val)
            cell.font = font_code if col_idx == 1 else (font_bold if col_idx == 4 else font_regular)
            cell.border = thin_border
            cell.fill = fill
            cell.alignment = align_center if col_idx in [1, 4] else align_left

    # =============================================================
    # SHEET 5: Guide_For_Polymator
    # =============================================================
    ws5 = wb.create_sheet(title="Guide_For_Polymator")
    ws5.views.sheetView[0].showGridLines = True

    ws5.merge_cells("A1:G1")
    ws5["A1"] = "Polymator Implementation & Excel Remaking Guide"
    ws5["A1"].font = font_title
    ws5["A1"].fill = fill_navy
    ws5["A1"].alignment = align_center
    ws5.row_dimensions[1].height = 36

    ws5.merge_cells("A2:G2")
    ws5["A2"] = "Clear instructions on how Polymator must build, structure, and link the final Excel sheet"
    ws5["A2"].font = font_subtitle
    ws5["A2"].fill = fill_blue_hdr
    ws5["A2"].alignment = align_center
    ws5.row_dimensions[2].height = 20

    guide_sections = [
        ("1. WHY THE OLD EXCEL CAUSED CONFUSION", [
            "• The previous master Excel (J886) merged 15 different domains into one sheet: thermodynamic calculations in Watts, 27 electrical circuits, TCMS cab switches, and 100+ hardware test points.",
            "• Excel Sir looked at that 1,500-row sheet and did not know which rows belonged to Unity vs LabVIEW internal math.",
            "• He specifically created the 10-item 'TO UNITY' list to declare: 'This is the bare minimum we need for the 3D visual software.'",
            "• RULE FOR POLYMATOR: Keep the Unity interface sheet 100% CLEAN. Never paste LabVIEW internal thermodynamic formulas or cab switch tables into the Unity exchange table."
        ]),
        ("2. STRUCTURE OF THE UNITY INTERFACE SHEETS", [
            "• TABLE A: LabVIEW_to_Unity (Controls - UDP Port 9000): Exactly 41 structured control rows divided into 6 logical categories:",
            "    - Category 1: Airflow Arrow Speeds (6 variables: BlueArrowOut, Fan1, Fan2, GreenArrowIN1, GreenArrowIN2, YellowArrowIn) [Req 1]",
            "    - Category 2: Airflow Arrow Color Lerp (5 variables: RedArrow_Inlet/Outlet, BlueArrow_Outlet, Green/Yellow Inlet) [Req 2]",
            "    - Category 3: Damper BlendShapes (3 variables: FADL, FADR, RAD) [Req 4, 5, 6]",
            "    - Category 4: App & Casing State (2 variables: DS3D_OFF, ShowHideObjects) [Req 3, 9]",
            "    - Category 5: Fault Trigger & Reset (3 variables: Fault_Code_Trigger, Fault_Cause_Trigger, Fault_Command_Action) [Req 7, 8]",
            "    - Category 6: Component Highlights (22 variables: 4 Compressors, 3 Dampers, 7 Sensors, 4 Fans/Blowers, 4 Valves/Heater) [Req 10]",
            "• TABLE B: Unity_to_LabVIEW (Telemetry - UDP Port 5005): Exactly 23 structured telemetry rows:",
            "    - Category 1: Procedural Scoring (9 variables: Active_Fault_Code, Current_Step_Number, Total_Steps, Wrong_Attempts, Step_Accomplished, Fault_Completed, Fault_Submit_Status, Steps_Attempts_Summary, Telemetry_Event)",
            "    - Category 2: Discrete Status Bits (14 booleans: Fault_101_Status to Fault_114_Status)"
        ]),
        ("3. ANSWERS TO TECHNICAL SPECIFICATION QUESTIONS", [
            "• Fresh Air Damper: Has BOTH BlendShape animation (FADL & FADR from 0% closed to 100% open) AND highlights for repair procedures.",
            "• Return Air Damper (RAD): Controls return air blade BlendShape (Normal = 100% open; Smoke Mode = 0% closed).",
            "• ShowHideObjects: Toggles outer metal cover in Unity. Value 0 hides casing so student sees inside; Value 1 restores the casing.",
            "• All Component Highlights: Covers all 4 compressors, 3 dampers, 7 sensors, 4 fans/blowers, and 4 valves/heater elements."
        ]),
        ("4. WHY TELEMETRY VARIABLES ARE ESSENTIAL FOR SCORING", [
            "• Current_Step_Number & Total_Steps: LabVIEW needs them to render the progress bar and know when the student reaches the final step.",
            "• Wrong_Attempts: Increments whenever the student selects the wrong tool. LabVIEW uses this count to deduct marks in Exam Mode.",
            "• Step_Accomplished & Fault_Completed: Pulses and flags that tell LabVIEW to advance timers and generate the final grade report.",
            "• Telemetry_Event: String field ('step_completed', 'fault_completed') used by LabVIEW's native 'Unflatten From JSON' block."
        ])
    ]

    curr_r = 4
    for section_title, bullets in guide_sections:
        ws5.merge_cells(f"A{curr_r}:G{curr_r}")
        ws5[f"A{curr_r}"] = section_title
        ws5[f"A{curr_r}"].font = Font(name="Segoe UI", size=11, bold=True, color="1F4E79")
        ws5[f"A{curr_r}"].fill = fill_cat_blue
        ws5.row_dimensions[curr_r].height = 22
        curr_r += 1

        for b in bullets:
            ws5.merge_cells(f"A{curr_r}:G{curr_r}")
            ws5[f"A{curr_r}"] = b
            ws5[f"A{curr_r}"].font = font_regular
            ws5.row_dimensions[curr_r].height = 20
            curr_r += 1
        curr_r += 1

    # Auto-adjust column widths across all sheets
    for ws in [ws1, ws2, ws3, ws4, ws5]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # ignore merged cells spanning across
                if cell.coordinate in ws.merged_cells:
                    continue
                val_str = str(cell.value or '')
                if val_str:
                    max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Set specific balanced widths for critical sheets
    ws1.column_dimensions["A"].width = 8
    ws1.column_dimensions["B"].width = 30
    ws1.column_dimensions["C"].width = 18
    ws1.column_dimensions["D"].width = 36
    ws1.column_dimensions["E"].width = 24
    ws1.column_dimensions["F"].width = 34
    ws1.column_dimensions["G"].width = 38

    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 30
    ws2.column_dimensions["C"].width = 32
    ws2.column_dimensions["D"].width = 14
    ws2.column_dimensions["E"].width = 18
    ws2.column_dimensions["F"].width = 10
    ws2.column_dimensions["G"].width = 10
    ws2.column_dimensions["H"].width = 46

    ws3.column_dimensions["A"].width = 8
    ws3.column_dimensions["B"].width = 26
    ws3.column_dimensions["C"].width = 32
    ws3.column_dimensions["D"].width = 14
    ws3.column_dimensions["E"].width = 20
    ws3.column_dimensions["F"].width = 10
    ws3.column_dimensions["G"].width = 34
    ws3.column_dimensions["H"].width = 46

    ws4.column_dimensions["A"].width = 12
    ws4.column_dimensions["B"].width = 32
    ws4.column_dimensions["C"].width = 28
    ws4.column_dimensions["D"].width = 14
    ws4.column_dimensions["E"].width = 34
    ws4.column_dimensions["F"].width = 34
    ws4.column_dimensions["G"].width = 36

    print(f"Saving workbook to {output_path}...")
    wb.save(output_path)
    print("Workbook created and saved successfully!")

if __name__ == "__main__":
    build_excel()
